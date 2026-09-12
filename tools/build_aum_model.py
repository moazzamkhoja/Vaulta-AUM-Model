# -*- coding: utf-8 -*-
"""
Vaulta_AUM_Model_v1.xlsx
AUM accumulation model: customers x adoption x contribution %, compounding
through contributions AND investment returns, charged at the AUM fee.
Tests whether AUM fee revenue overtakes T-bill NIM in Years 3-5.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = r"C:\Projects\Vaulta-AUM-Model"
OUT = os.path.join(REPO, "Vaulta_AUM_Model_v1.xlsx")
NQ = 20
FIRST_Q_COL = 3

NAVY, BLUE, LTBLUE, GREY = "1F4E79", "2E75B6", "BDD7EE", "D9E1F2"
IN_BLUE = "0000FF"
F_IN   = Font(color=IN_BLUE, size=10)
F_FM   = Font(color="000000", size=10)
F_FMB  = Font(color="000000", size=10, bold=True)
F_TITLE= Font(color="FFFFFF", size=14, bold=True)
F_SEC  = Font(color="FFFFFF", size=10, bold=True)
F_HDR  = Font(color="FFFFFF", size=9, bold=True)
F_NOTE = Font(color="808080", size=8, italic=True)
F_RED  = Font(color="C00000", size=9, bold=True)
F_SRC  = Font(color="2E75B6", size=8)

FILL_TITLE = PatternFill("solid", fgColor=NAVY)
FILL_SEC   = PatternFill("solid", fgColor=BLUE)
FILL_HDR   = PatternFill("solid", fgColor=NAVY)
FILL_IN    = PatternFill("solid", fgColor="FFF8E7")
FILL_LT    = PatternFill("solid", fgColor=LTBLUE)
FILL_WARN  = PatternFill("solid", fgColor="FDE9E9")

THIN = Side(style="thin", color="BFBFBF")
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOPL = Border(top=Side(style="thin", color=NAVY))

M0, M2, PCT1, PCT2, NUM, XF = '$#,##0', '$#,##0.00', '0.0%', '0.00%', '#,##0', '0.00"x"'

def title(ws, row, text, span=10):
    ws.cell(row, 1, text).font = F_TITLE
    for c in range(1, span + 1): ws.cell(row, c).fill = FILL_TITLE
    ws.row_dimensions[row].height = 22

def section(ws, row, text, span=10):
    ws.cell(row, 1, text).font = F_SEC
    for c in range(1, span + 1): ws.cell(row, c).fill = FILL_SEC
    ws.row_dimensions[row].height = 16

def note(ws, row, text, col=1):
    ws.cell(row, col, text).font = F_NOTE

def label(ws, row, text, col=2, bold=False):
    c = ws.cell(row, col, text); c.font = F_FMB if bold else F_FM; return c

def put_in(ws, row, col, value, fmt=None):
    c = ws.cell(row, col, value); c.font = F_IN; c.fill = FILL_IN; c.border = BOX
    if fmt: c.number_format = fmt
    return c

def put_fm(ws, row, col, formula, fmt=None, bold=False):
    c = ws.cell(row, col, formula); c.font = F_FMB if bold else F_FM
    if fmt: c.number_format = fmt
    return c

def hdr(ws, row, col, text, align="center"):
    c = ws.cell(row, col, text); c.font = F_HDR; c.fill = FILL_HDR
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    c.border = BOX; return c

def qcol(q): return get_column_letter(FIRST_Q_COL + q - 1)
def widths(ws, spec):
    for k, v in spec.items(): ws.column_dimensions[k].width = v

wb = Workbook()
ws_cov = wb.active; ws_cov.title = "Cover"
ws_res = wb.create_sheet("Research & Sources")
ws_asm = wb.create_sheet("Assumptions")
ws_con = wb.create_sheet("Consumer Model")
ws_eng = wb.create_sheet("AUM Engine")
ws_rev = wb.create_sheet("Revenue Comparison")
ws_cx  = wb.create_sheet("Crossover Analysis")
ws_ue  = wb.create_sheet("Segment Unit Economics")
ws_sen = wb.create_sheet("Sensitivity")

S_ASM, S_CON = "Assumptions", "'Consumer Model'"
S_ENG, S_REV = "'AUM Engine'", "'Revenue Comparison'"
S_CX, S_UE   = "'Crossover Analysis'", "'Segment Unit Economics'"
R = {}

SEGS = [
    ("1 — Underbanked Active",      40.00,  900, 0.0325, 0.08, 0.010, 0.25),
    ("2 — Dissatisfied Banked",     57.50, 2500, 0.0225, 0.25, 0.020, 0.60),
    ("3 — Primary Bank Switcher",   70.00, 4250, 0.0125, 0.40, 0.025, 1.00),
    ("4 — Crypto-Curious Investor", 47.50, 6000, 0.0150, 0.70, 0.035, 2.00),
]
MIX = [
    ("1 — Underbanked Active",      [0.60, 0.475, 0.35, 0.35, 0.35]),
    ("2 — Dissatisfied Banked",     [0.30, 0.375, 0.45, 0.45, 0.45]),
    ("3 — Primary Bank Switcher",   [0.00, 0.075, 0.15, 0.15, 0.15]),
    ("4 — Crypto-Curious Investor", [0.10, 0.075, 0.05, 0.05, 0.05]),
]
seg_names = [s[0] for s in SEGS]

# ═══════════════════════════════════════════════════ RESEARCH & SOURCES
ws = ws_res
widths(ws, {'A': 3, 'B': 34, 'C': 14, 'D': 52, 'E': 60})
title(ws, 1, "  RESEARCH & SOURCES  —  basis for every default on the Assumptions sheet", 5)
note(ws, 2, "  Researched September 2026. Where a source gave a range, the model default sits inside it and the range is stated.")
r = 4
for col, txt in [(2, "Assumption"), (3, "Default"), (4, "Basis"), (5, "Source")]:
    hdr(ws, r, col, txt, "left")
r += 1
RESEARCH = [
 ("AUM fee", "0.50%/yr",
  "Betterment Basic and Wealthfront both charge 0.25%. Acorns charges a flat $3/mo, "
  "which on its ~$2,142 average balance is ~1.7% and ~2.5% on a $1,425 balance. "
  "0.50% sits between the two and undercuts Acorns by ~3x for small accounts.",
  "unbiased.com robo-advisor comparison 2026; acorns.com/learn fee comparison"),
 ("Avg investment balance\n(reference point)", "$2,142",
  "Acorns: $30bn AUM across 14m registered users (Jul 2026); $1,876 in 2025, $1,439 in 2024. "
  "Closest mass-market analogue to Vaulta's segments. Balance per FUNDED account is higher "
  "since registered exceeds funded.",
  "investingintheweb.com/brokers/acorn-statistics"),
 ("Monthly contribution", "~$50/mo\n(seg 2)",
  "Acorns reports most customers investing roughly $50-60/month into core accounts. "
  "Modelled here as a % of wallet balance so it scales with segment wealth.",
  "investingintheweb.com/brokers/acorn-statistics"),
 ("Contribution % of balance", "1.0 - 3.5%/mo",
  "US personal saving rate was 2.7-3.0% of income in mid-2026. Top-quintile earners save "
  "15-25%; bottom-quintile saving is negative. Segment rates are scaled to that spread and "
  "calibrated so segment 2 lands at the Acorns $50/mo observation.",
  "tradingeconomics.com US personal saving rate; Federal Reserve SCF via Bankrate"),
 ("Expected investment return", "5.5%/yr",
  "Vanguard's Dec-2025 10-year US equity forecast is 3.9-5.9%; BlackRock's was just over 5% "
  "in Sep-2025. Both are well below the historical ~10% and reflect current valuations. "
  "5.5% is the upper-middle of that consensus.",
  "corporate.vanguard.com VEMO 2026; Morningstar 2026 expert forecast survey"),
 ("Affluent-investor ceiling\n(reference)", "$65-68K",
  "Betterment ~$65bn across 1m+ clients; Wealthfront $95bn across 1.4m+. This is the "
  "primary-brokerage tier, roughly 30x the Acorns balance — the gap that decides whether "
  "AUM can overtake NIM.",
  "sacra.com robo-advisor research; investingintheweb Betterment statistics"),
 ("Product is real", "Live Aug 2026",
  "Dinari launched 724 tokenized US stocks including the full S&P 500 for US self-custody "
  "wallets, settled in USDC across 4 chains, each dShare backed 1:1 in regulated custody. "
  "The Vaulta Invest premise is available today, not speculative.",
  "coindesk.com 2026-08-04; theblock.co/post/410588"),
 ("Market backdrop", "$5.5T by 2030",
  "Citigroup projects the tokenized securities market reaches $5.5tn by 2030. Ondo has an "
  "SEC-aligned framework using iShares Core S&P 500 as collateral, not yet open to US retail.",
  "coindesk.com tokenized equities coverage 2026"),
 ("Micro-investing growth", "+27% YoY",
  "Investment fintech apps grew users 27% year over year, led by micro-investing and "
  "fractional-share platforms. Micro-investing platform revenue $0.93bn in 2025 -> $1.85bn by 2029.",
  "globenewswire.com Micro Investing Platforms Market Report 2025-2029"),
 ("Redemption / leakage", "8%/yr",
  "No clean public disclosure for micro-investing redemption rates. 8% is a modelling "
  "judgement, not a researched figure — it is an input and is flagged as the weakest default "
  "in the workbook.",
  "MODELLING JUDGEMENT — not sourced"),
]
for a, d, b, s in RESEARCH:
    c = ws.cell(r, 2, a); c.font = F_FMB; c.alignment = Alignment(wrap_text=True, vertical="top")
    c = ws.cell(r, 3, d); c.font = F_IN;  c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
    c = ws.cell(r, 4, b); c.font = Font(size=9); c.alignment = Alignment(wrap_text=True, vertical="top")
    c = ws.cell(r, 5, s); c.font = F_SRC; c.alignment = Alignment(wrap_text=True, vertical="top")
    if s.startswith("MODELLING"):
        for cc in range(2, 6): ws.cell(r, cc).fill = FILL_WARN
    for cc in range(2, 6): ws.cell(r, cc).border = BOX
    ws.row_dimensions[r].height = 52
    r += 1

# ═══════════════════════════════════════════════════ ASSUMPTIONS
ws = ws_asm
widths(ws, {'A': 3, 'B': 44, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 34})
title(ws, 1, "  VAULTA AUM MODEL — ASSUMPTIONS", 8)
note(ws, 2, "  BLUE = input.  BLACK = formula.  Defaults are researched; see the Research & Sources sheet for the basis of each.")
r = 4

section(ws, r, "  1.  T-BILL SLEEVE  (the revenue AUM is being compared against)", 8); r += 1
label(ws, r, "T-bill annual rate");               R['tbill'] = r; put_in(ws, r, 3, 0.048, PCT2); r += 1
label(ws, r, "Vaulta sleeve % (kept as NIM)");    R['sleeve'] = r; put_in(ws, r, 3, 0.20, PCT1)
ws.cell(r, 8, "20/80 split — Vaulta keeps 20%").font = F_NOTE; r += 1
label(ws, r, "Consumer reward % (passed through)"); R['reward'] = r
put_fm(ws, r, 3, f"=1-C{R['sleeve']}", PCT1); r += 1
label(ws, r, "Consumer effective reward rate");   R['effrwd'] = r
put_fm(ws, r, 3, f"=C{R['tbill']}*C{R['reward']}", PCT2)
ws.cell(r, 8, "vs Chime 3.75% APY (needs $3k DD)").font = F_NOTE; r += 1
label(ws, r, "Vaulta NIM rate on balances");      R['nimrate'] = r
put_fm(ws, r, 3, f"=C{R['tbill']}*C{R['sleeve']}", PCT2); r += 2

section(ws, r, "  2.  VAULTA INVEST  —  AUM ENGINE", 8); r += 1
label(ws, r, "AUM fee (annual)");                 R['aumfee'] = r; put_in(ws, r, 3, 0.005, PCT2); r += 1
label(ws, r, "Expected investment return (annual)"); R['mktret'] = r; put_in(ws, r, 3, 0.055, PCT1)
ws.cell(r, 8, "Vanguard 3.9-5.9%, BlackRock ~5%").font = F_NOTE; r += 1
label(ws, r, "Annual redemption / leakage rate"); R['redeem'] = r; put_in(ws, r, 3, 0.08, PCT1)
ws.cell(r, 8, "modelling judgement — weakest default").font = F_NOTE; r += 1
for i, y in enumerate(range(1, 6)): hdr(ws, r, 3 + i, f"Y{y}")
label(ws, r, ""); r += 1
label(ws, r, "Product availability ramp");        R['ramp'] = r
for i, v in enumerate([0.00, 0.45, 0.70, 0.90, 1.00]): put_in(ws, r, 3 + i, v, PCT1)
ws.cell(r, 8, "scales each segment's ceiling").font = F_NOTE; r += 1
note(ws, r, "  One adoption mechanism only: effective adoption = segment ceiling x availability ramp. Avoids the two-schedule conflict in the Pay-Model."); r += 2

section(ws, r, "  3.  SEGMENT LEVERS", 8); r += 1
for col, txt in [(2, "Segment"), (3, "CAC"), (4, "Wallet balance"), (5, "Monthly churn"),
                 (6, "AUM adoption ceiling"), (7, "Monthly contribution\n(% of balance)"),
                 (8, "Initial rollover at signup\n(x wallet balance)")]:
    hdr(ws, r, col, txt)
ws.row_dimensions[r].height = 30
r += 1
R['seg0'] = r
for nm, cac, bal, chn, adopt, contrib, roll in SEGS:
    label(ws, r, nm)
    put_in(ws, r, 3, cac, M2); put_in(ws, r, 4, bal, M0); put_in(ws, r, 5, chn, PCT2)
    put_in(ws, r, 6, adopt, PCT1); put_in(ws, r, 7, contrib, PCT1); put_in(ws, r, 8, roll, XF)
    r += 1
R['seg1'] = r - 1
r += 1
label(ws, r, "memo: implied monthly contribution ($)", bold=True); r += 1
R['contrib0'] = r
for i in range(4):
    label(ws, r, "   " + seg_names[i])
    put_fm(ws, r, 3, f"=D{R['seg0']+i}*G{R['seg0']+i}", M2)
    r += 1
R['contrib1'] = r - 1
note(ws, r, "  Segment 2 lands near the ~$50/month Acorns observation. Initial rollover is the lump sum a new investor transfers in at signup, as a multiple of wallet balance."); r += 2

section(ws, r, "  4.  SEGMENT MIX BY YEAR", 8); r += 1
for i, y in enumerate(range(1, 6)): hdr(ws, r, 3 + i, f"Y{y}")
label(ws, r, "Segment"); r += 1
R['mix0'] = r
for nm, vals in MIX:
    label(ws, r, nm)
    for i, v in enumerate(vals): put_in(ws, r, 3 + i, v, PCT1)
    r += 1
R['mix1'] = r - 1
label(ws, r, "Total", bold=True)
for i in range(5):
    cl = get_column_letter(3 + i)
    put_fm(ws, r, 3 + i, f"=SUM({cl}{R['mix0']}:{cl}{R['mix1']})", PCT1, bold=True)
r += 2

section(ws, r, "  5.  CONSUMER GROWTH", 8); r += 1
label(ws, r, "Starting active consumers (Q0)");  R['q0'] = r; put_in(ws, r, 3, 1000, NUM); r += 1
for i, y in enumerate(range(1, 6)): hdr(ws, r, 3 + i, f"Y{y}")
label(ws, r, ""); r += 1
label(ws, r, "Year-end active consumers");       R['yearend'] = r
for i, v in enumerate([10000, 35000, 100000, 250000, 500000]): put_in(ws, r, 3 + i, v, NUM)
r += 2

section(ws, r, "  6.  TRANSACTION FEES & COSTS  (context for unit economics)", 8); r += 1
label(ws, r, "Monthly spend as % of balance");   R['spend'] = r; put_in(ws, r, 3, 0.20, PCT1); r += 1
label(ws, r, "Share of spend on QR / Solana Pay"); R['qrs'] = r; put_in(ws, r, 3, 0.40, PCT1); r += 1
label(ws, r, "Share of spend on Rain card");     R['cds'] = r; put_fm(ws, r, 3, f"=1-C{R['qrs']}", PCT1); r += 1
label(ws, r, "QR merchant fee");                 R['qrf'] = r; put_in(ws, r, 3, 0.005, PCT2); r += 1
label(ws, r, "Rain net interchange");            R['rin'] = r; put_in(ws, r, 3, 0.0035, PCT2); r += 1
label(ws, r, "Rain per-transaction cost");       R['rct'] = r; put_in(ws, r, 3, 0.15, M2); r += 1
label(ws, r, "Average transaction size");        R['tkt'] = r; put_in(ws, r, 3, 45.00, M2); r += 1
label(ws, r, "Infra per consumer / month");      R['infra'] = r; put_in(ws, r, 3, 0.50, M2); r += 2

section(ws, r, "  7.  LTV DISCOUNTING", 8); r += 1
label(ws, r, "WACC");                            R['wacc'] = r; put_in(ws, r, 3, 0.15, PCT1); r += 1
label(ws, r, "Monthly discount rate");           R['mdisc'] = r
put_fm(ws, r, 3, f"=(1+C{R['wacc']})^(1/12)-1", '0.000%'); r += 1
label(ws, r, "LTV horizon (months)");            R['horizon'] = r; put_in(ws, r, 3, 60, NUM); r += 1

def A(key, col=3):
    return f"{S_ASM}!${get_column_letter(col)}${R[key]}"
def Arow(key, c0=3, c1=7):
    return f"{S_ASM}!${get_column_letter(c0)}${R[key]}:${get_column_letter(c1)}${R[key]}"
def Ablk(k0, k1, c0=3, c1=7):
    return f"{S_ASM}!${get_column_letter(c0)}${R[k0]}:${get_column_letter(c1)}${R[k1]}"

TB, SLV, RWD = A('tbill'), A('sleeve'), A('reward')
AUMF, MKTR, REDM = A('aumfee'), A('mktret'), A('redeem')
SPND, QRS, CDS = A('spend'), A('qrs'), A('cds')
QRF, RIN, RCT, TKT, INFR = A('qrf'), A('rin'), A('rct'), A('tkt'), A('infra')
MDSC, HORZ = A('mdisc'), A('horizon')
MIXBLK = Ablk('mix0', 'mix1')
SEG_CAC = f"{S_ASM}!$C${R['seg0']}:$C${R['seg1']}"
SEG_BAL = f"{S_ASM}!$D${R['seg0']}:$D${R['seg1']}"
SEG_CHN = f"{S_ASM}!$E${R['seg0']}:$E${R['seg1']}"

QH = 5
def qtr_header(ws, base_year=2027):
    hdr(ws, QH, 2, "Quarter", "left")
    for q in range(1, NQ + 1):
        c = ws.cell(QH, FIRST_Q_COL + q - 1, q)
        c.font = F_HDR; c.fill = FILL_HDR
        c.alignment = Alignment(horizontal="center"); c.border = BOX
    label(ws, QH + 1, "Year"); label(ws, QH + 2, "Period")
    for q in range(1, NQ + 1):
        put_fm(ws, QH + 1, FIRST_Q_COL + q - 1, f"=INT(({qcol(q)}{QH}-1)/4)+1", NUM)
        y = base_year + (q - 1) // 4
        p = ws.cell(QH + 2, FIRST_Q_COL + q - 1, f"{y} Q{(q-1)%4+1}")
        p.font = F_NOTE; p.alignment = Alignment(horizontal="center")

def qwidths(ws, lw=42):
    widths(ws, {'A': 3, 'B': lw})
    for q in range(1, NQ + 1): ws.column_dimensions[qcol(q)].width = 12

def yr(q): return f"{S_CON}!{qcol(q)}${QH+1}"

def qrow(ws, row, rlabel, fn, fmt, bold=False, fill=None, indent=False):
    label(ws, row, ("   " if indent else "") + rlabel, bold=bold)
    for q in range(1, NQ + 1):
        c = put_fm(ws, row, FIRST_Q_COL + q - 1, fn(q), fmt, bold=bold)
        if fill: c.fill = fill
    return row

# ═══════════════════════════════════════════════════ CONSUMER MODEL
ws = ws_con
qwidths(ws)
title(ws, 1, "  CONSUMER MODEL  —  quarterly ramp by segment", 8)
qtr_header(ws)
r = QH + 4
section(ws, r, "  ACTIVE CONSUMERS", 8); r += 1
c_tgt = qrow(ws, r, "Year-end target", lambda q: f"=INDEX({Arow('yearend')},1,{yr(q)})", NUM, indent=True); r += 1
c_pri = qrow(ws, r, "Prior year-end base",
             lambda q: f"=IF({yr(q)}=1,{A('q0')},INDEX({Arow('yearend')},1,{yr(q)}-1))", NUM, indent=True); r += 1
c_eop = qrow(ws, r, "Active consumers (end of period)",
             lambda q: f"={qcol(q)}{c_pri}*({qcol(q)}{c_tgt}/{qcol(q)}{c_pri})^(({qcol(q)}{QH}-({yr(q)}-1)*4)/4)",
             NUM, bold=True, fill=FILL_LT); r += 1
c_bop = qrow(ws, r, "Active consumers (beginning)",
             lambda q: f"={A('q0')}" if q == 1 else f"={qcol(q-1)}{c_eop}", NUM, indent=True); r += 1
c_avg = qrow(ws, r, "Active consumers (average)",
             lambda q: f"=({qcol(q)}{c_bop}+{qcol(q)}{c_eop})/2", NUM); r += 1
c_mch = qrow(ws, r, "Blended monthly churn",
             lambda q: f"=SUMPRODUCT({SEG_CHN},INDEX({MIXBLK},0,{yr(q)}))", PCT2); r += 1
c_qch = qrow(ws, r, "Quarterly churn rate",
             lambda q: f"=1-(1-{qcol(q)}{c_mch})^3", PCT1, indent=True); r += 1
c_cac = qrow(ws, r, "Blended CAC",
             lambda q: f"=SUMPRODUCT({SEG_CAC},INDEX({MIXBLK},0,{yr(q)}))", M2); r += 1
c_bal = qrow(ws, r, "Blended wallet balance",
             lambda q: f"=SUMPRODUCT({SEG_BAL},INDEX({MIXBLK},0,{yr(q)}))", M0); r += 1
c_tba = qrow(ws, r, "TOTAL WALLET BALANCES (avg)",
             lambda q: f"={qcol(q)}{c_avg}*{qcol(q)}{c_bal}", M0, bold=True, fill=FILL_LT); r += 2

section(ws, r, "  ACTIVE CONSUMERS BY SEGMENT (end of period)", 8); r += 1
c_seg = []
for i, nm in enumerate(seg_names):
    c_seg.append(qrow(ws, r, nm, lambda q, i=i: f"={qcol(q)}{c_eop}*INDEX({MIXBLK},{i+1},{yr(q)})", NUM, indent=True)); r += 1
r += 1
section(ws, r, "  ACTIVE CONSUMERS BY SEGMENT (average)", 8); r += 1
c_sgav = []
for i, nm in enumerate(seg_names):
    c_sgav.append(qrow(ws, r, nm, lambda q, i=i: f"={qcol(q)}{c_avg}*INDEX({MIXBLK},{i+1},{yr(q)})", NUM, indent=True)); r += 1

# ═══════════════════════════════════════════════════ AUM ENGINE
ws = ws_eng
qwidths(ws, 46)
title(ws, 1, "  AUM ENGINE  —  quarterly roll-forward by segment", 8)
note(ws, 2, "  Closing AUM = opening + initial rollovers + contributions - redemptions - churn leakage + market return.  Fee is charged on average AUM.")
qtr_header(ws)
r = QH + 4

SEG_ROWS = {}
for i, nm in enumerate(seg_names):
    arow = R['seg0'] + i
    section(ws, r, f"  SEGMENT {nm}", 8); r += 1
    s_usr = qrow(ws, r, "Invest users (end of period)",
                 lambda q, i=i, a=arow: f"={S_CON}!{qcol(q)}{c_seg[i]}*{S_ASM}!$F${a}*INDEX({Arow('ramp')},1,{yr(q)})",
                 NUM, indent=True); r += 1
    s_new = qrow(ws, r, "New invest users",
                 lambda q, u=s_usr: (f"={qcol(q)}{u}" if q == 1 else f"=MAX(0,{qcol(q)}{u}-{qcol(q-1)}{u})"),
                 NUM, indent=True); r += 1
    s_op = qrow(ws, r, "Opening AUM", lambda q: "=0", M0, indent=True); r += 1
    s_roll = qrow(ws, r, "+ Initial rollovers",
                  lambda q, n=s_new, a=arow: f"={qcol(q)}{n}*{S_ASM}!$D${a}*{S_ASM}!$H${a}", M0, indent=True); r += 1
    s_con = qrow(ws, r, "+ Contributions",
                 lambda q, u=s_usr, a=arow: f"={qcol(q)}{u}*{S_ASM}!$D${a}*{S_ASM}!$G${a}*3", M0, indent=True); r += 1
    s_red = qrow(ws, r, "- Redemptions",
                 lambda q, o=s_op: f"=-{qcol(q)}{o}*{REDM}/4", M0, indent=True); r += 1
    s_chn = qrow(ws, r, "- Churn leakage",
                 lambda q, o=s_op, a=arow: f"=-{qcol(q)}{o}*(1-(1-{S_ASM}!$E${a})^3)", M0, indent=True); r += 1
    s_ret = qrow(ws, r, "+ Market return",
                 lambda q, o=s_op, roll=s_roll, con=s_con: f"=({qcol(q)}{o}+({qcol(q)}{roll}+{qcol(q)}{con})/2)*((1+{MKTR})^(1/4)-1)",
                 M0, indent=True); r += 1
    s_cl = qrow(ws, r, "CLOSING AUM",
                lambda q, o=s_op, a=s_roll, b=s_con, c=s_red, d=s_chn, e=s_ret:
                    f"={qcol(q)}{o}+{qcol(q)}{a}+{qcol(q)}{b}+{qcol(q)}{c}+{qcol(q)}{d}+{qcol(q)}{e}",
                M0, bold=True, fill=FILL_LT); r += 1
    for q in range(1, NQ + 1):
        ws.cell(s_op, FIRST_Q_COL + q - 1).value = "=0" if q == 1 else f"={qcol(q-1)}{s_cl}"
    s_avg = qrow(ws, r, "Average AUM",
                 lambda q, o=s_op, c=s_cl: f"=({qcol(q)}{o}+{qcol(q)}{c})/2", M0, indent=True); r += 1
    s_per = qrow(ws, r, "AUM per invest user",
                 lambda q, c=s_cl, u=s_usr: f"=IF({qcol(q)}{u}=0,0,{qcol(q)}{c}/{qcol(q)}{u})", M0, indent=True); r += 2
    SEG_ROWS[i] = dict(usr=s_usr, new=s_new, op=s_op, roll=s_roll, con=s_con,
                       red=s_red, chn=s_chn, ret=s_ret, cl=s_cl, avg=s_avg, per=s_per)

section(ws, r, "  TOTAL — ALL SEGMENTS", 8); r += 1
def tot(key):
    return lambda q: "=" + "+".join(f"{qcol(q)}{SEG_ROWS[i][key]}" for i in range(4))
t_usr = qrow(ws, r, "Total invest users", tot('usr'), NUM, bold=True); r += 1
t_roll= qrow(ws, r, "Initial rollovers", tot('roll'), M0, indent=True); r += 1
t_con = qrow(ws, r, "Contributions", tot('con'), M0, indent=True); r += 1
t_red = qrow(ws, r, "Redemptions", tot('red'), M0, indent=True); r += 1
t_chn = qrow(ws, r, "Churn leakage", tot('chn'), M0, indent=True); r += 1
t_ret = qrow(ws, r, "Market return", tot('ret'), M0, indent=True); r += 1
t_cl  = qrow(ws, r, "TOTAL CLOSING AUM", tot('cl'), M0, bold=True, fill=FILL_LT); r += 1
for q in range(1, NQ + 1): ws.cell(t_cl, FIRST_Q_COL + q - 1).border = TOPL
t_avg = qrow(ws, r, "Total average AUM", tot('avg'), M0); r += 1
t_per = qrow(ws, r, "AUM per invest user",
             lambda q: f"=IF({qcol(q)}{t_usr}=0,0,{qcol(q)}{t_cl}/{qcol(q)}{t_usr})", M0, bold=True); r += 1
t_adp = qrow(ws, r, "Effective adoption (% of actives)",
             lambda q: f"={qcol(q)}{t_usr}/{S_CON}!{qcol(q)}{c_eop}", PCT1); r += 2

section(ws, r, "  AUM FEE REVENUE", 8); r += 1
t_fee = qrow(ws, r, "AUM fee revenue (quarter)",
             lambda q: f"={qcol(q)}{t_avg}*{AUMF}/4", M0, bold=True, fill=FILL_LT); r += 1
t_arr = qrow(ws, r, "AUM fee (annual run-rate)",
             lambda q: f"={qcol(q)}{t_cl}*{AUMF}", M0); r += 1

# ═══════════════════════════════════════════════════ REVENUE COMPARISON
ws = ws_rev
qwidths(ws, 44)
title(ws, 1, "  REVENUE COMPARISON  —  does AUM overtake the T-bill sleeve?", 8)
note(ws, 2, "  Compared on what Vaulta KEEPS. T-bill contributes only the 20% sleeve; the 80% consumer reward is passed through and is not Vaulta revenue.")
qtr_header(ws)
r = QH + 4

section(ws, r, "  VAULTA RETAINED REVENUE BY LINE", 8); r += 1
x_bal = qrow(ws, r, "Total wallet balances (avg)",
             lambda q: f"={S_CON}!{qcol(q)}{c_tba}", M0, indent=True); r += 1
x_tbg = qrow(ws, r, "T-bill gross income",
             lambda q: f"={qcol(q)}{x_bal}*{TB}/4", M0, indent=True); r += 1
x_rwd = qrow(ws, r, "   less consumer reward (80%)",
             lambda q: f"=-{qcol(q)}{x_tbg}*{RWD}", M0, indent=True); r += 1
x_nim = qrow(ws, r, "T-BILL NIM (Vaulta keeps)",
             lambda q: f"={qcol(q)}{x_tbg}*{SLV}", M0, bold=True, fill=FILL_LT); r += 1
x_spd = qrow(ws, r, "   memo: consumer spend",
             lambda q: f"={qcol(q)}{x_bal}*{SPND}*3", M0, indent=True); r += 1
x_qrf = qrow(ws, r, "QR merchant fees",
             lambda q: f"={qcol(q)}{x_spd}*{QRS}*{QRF}", M0, indent=True); r += 1
x_rif = qrow(ws, r, "Rain interchange (net of txn cost)",
             lambda q: f"={qcol(q)}{x_spd}*{CDS}*{RIN}-{qcol(q)}{x_spd}*{CDS}/{TKT}*{RCT}", M0, indent=True); r += 1
x_txn = qrow(ws, r, "Transaction fees (QR + card net)",
             lambda q: f"={qcol(q)}{x_qrf}+{qcol(q)}{x_rif}", M0, bold=True); r += 1
x_aum = qrow(ws, r, "AUM FEE",
             lambda q: f"={S_ENG}!{qcol(q)}{t_fee}", M0, bold=True, fill=FILL_LT); r += 1
x_tot = qrow(ws, r, "TOTAL VAULTA REVENUE (retained)",
             lambda q: f"={qcol(q)}{x_nim}+{qcol(q)}{x_txn}+{qcol(q)}{x_aum}",
             M0, bold=True, fill=FILL_LT); r += 1
for q in range(1, NQ + 1): ws.cell(x_tot, FIRST_Q_COL + q - 1).border = TOPL
r += 1

section(ws, r, "  MIX", 8); r += 1
x_m1 = qrow(ws, r, "T-bill NIM %", lambda q: f"={qcol(q)}{x_nim}/{qcol(q)}{x_tot}", PCT1, indent=True); r += 1
x_m2 = qrow(ws, r, "Transaction fees %",
            lambda q: f"={qcol(q)}{x_txn}/{qcol(q)}{x_tot}", PCT1, indent=True); r += 1
x_m3 = qrow(ws, r, "AUM fee %", lambda q: f"={qcol(q)}{x_aum}/{qcol(q)}{x_tot}", PCT1, bold=True); r += 2

section(ws, r, "  THE TEST  —  AUM vs NIM", 8); r += 1
x_rat = qrow(ws, r, "AUM fee / T-bill NIM",
             lambda q: f"=IF({qcol(q)}{x_nim}=0,0,{qcol(q)}{x_aum}/{qcol(q)}{x_nim})", XF, bold=True, fill=FILL_LT); r += 1
x_ovr = qrow(ws, r, "AUM overtakes NIM?",
             lambda q: f'=IF({qcol(q)}{x_aum}>{qcol(q)}{x_nim},"YES","no")', None, indent=True); r += 2

section(ws, r, "  ANNUAL SUMMARY", 8); r += 1
YRROW = f"{S_CON}!$C${QH+1}:$V${QH+1}"
hdr(ws, r, 2, "Annual", "left")
for i in range(5): hdr(ws, r, 3 + i, f"Year {i+1}")
r += 1
def annrow(rlabel, srcrow, fmt, bold=False, fill=None):
    global r
    label(ws, r, rlabel, bold=bold)
    for i in range(5):
        c = put_fm(ws, r, 3 + i, f"=SUMIF({YRROW},{i+1},$C${srcrow}:$V${srcrow})", fmt, bold=bold)
        if fill: c.fill = fill
    this = r; r += 1; return this
an_nim = annrow("T-bill NIM (Vaulta keeps)", x_nim, M0, bold=True)
an_fee = annrow("Transaction fees", x_txn, M0)
an_aum = annrow("AUM fee", x_aum, M0, bold=True, fill=FILL_LT)
an_tot = annrow("Total retained revenue", x_tot, M0, bold=True)
label(ws, r, "AUM as % of retained revenue")
for i in range(5):
    cl = get_column_letter(3 + i)
    put_fm(ws, r, 3 + i, f"={cl}{an_aum}/{cl}{an_tot}", PCT1)
an_pct = r; r += 1
label(ws, r, "AUM / NIM ratio", bold=True)
for i in range(5):
    cl = get_column_letter(3 + i)
    c = put_fm(ws, r, 3 + i, f"={cl}{an_aum}/{cl}{an_nim}", XF, bold=True); c.fill = FILL_LT
an_rat = r; r += 1
label(ws, r, "AUM overtakes NIM?", bold=True)
for i in range(5):
    cl = get_column_letter(3 + i)
    put_fm(ws, r, 3 + i, f'=IF({cl}{an_aum}>{cl}{an_nim},"YES","no")', None, bold=True)
an_ovr = r; r += 1
label(ws, r, "memo: closing AUM (year end)")
for i in range(5):
    put_fm(ws, r, 3 + i, f"=INDEX({S_ENG}!$C${t_cl}:$V${t_cl},1,{(i+1)*4})", M0)
an_aumbal = r; r += 1
label(ws, r, "memo: total wallet balances (year end)")
for i in range(5):
    put_fm(ws, r, 3 + i, f"=INDEX({S_CON}!$C${c_tba}:$V${c_tba},1,{(i+1)*4})", M0)
an_wal = r; r += 1

# ═══════════════════════════════════════════════════ CROSSOVER ANALYSIS
ws = ws_cx
widths(ws, {'A': 3, 'B': 50, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': 16, 'H': 30})
title(ws, 1, "  CROSSOVER ANALYSIS  —  what would it take for AUM to overtake the T-bill sleeve?", 8)
note(ws, 2, "  AUM fee revenue = AUM x fee.  NIM revenue = balances x T-bill rate x sleeve.  Setting them equal gives a single required ratio of AUM to wallet balances.")
r = 4

section(ws, r, "  THE ALGEBRA", 8); r += 1
label(ws, r, "T-bill rate x sleeve % (NIM rate on balances)"); cx_nim = r
put_fm(ws, r, 3, f"={TB}*{SLV}", PCT2); r += 1
label(ws, r, "AUM fee rate"); cx_fee = r
put_fm(ws, r, 3, f"={AUMF}", PCT2); r += 1
label(ws, r, "REQUIRED AUM as multiple of wallet balances", bold=True); cx_req = r
c = put_fm(ws, r, 3, f"=C{cx_nim}/C{cx_fee}", XF, bold=True); c.fill = FILL_LT
ws.cell(r, 8, "total AUM must exceed this x total balances").font = F_NOTE; r += 2

section(ws, r, "  WHERE THE MODEL ACTUALLY LANDS", 8); r += 1
hdr(ws, r, 2, "Metric", "left")
for i in range(5): hdr(ws, r, 3 + i, f"Year {i+1}")
r += 1
label(ws, r, "Total AUM (year end)"); cx_aum = r
for i in range(5): put_fm(ws, r, 3 + i, f"={S_REV}!{get_column_letter(3+i)}${an_aumbal}", M0)
r += 1
label(ws, r, "Total wallet balances (year end)"); cx_bal = r
for i in range(5): put_fm(ws, r, 3 + i, f"={S_REV}!{get_column_letter(3+i)}${an_wal}", M0)
r += 1
label(ws, r, "Actual AUM / balances", bold=True); cx_act = r
for i in range(5):
    cl = get_column_letter(3 + i)
    c = put_fm(ws, r, 3 + i, f"={cl}{cx_aum}/{cl}{cx_bal}", XF, bold=True); c.fill = FILL_LT
r += 1
label(ws, r, "Required ratio (constant)"); cx_reqrow = r
for i in range(5): put_fm(ws, r, 3 + i, f"=$C${cx_req}", XF)
r += 1
label(ws, r, "Shortfall (x)", bold=True); cx_gap = r
for i in range(5):
    cl = get_column_letter(3 + i)
    c = put_fm(ws, r, 3 + i,
               f'=IF({cl}{cx_act}=0,"no AUM yet",{cl}{cx_reqrow}/{cl}{cx_act})', XF, bold=True)
    c.fill = FILL_WARN
r += 2

section(ws, r, "  WHAT WOULD CLOSE THE GAP", 8); r += 1
label(ws, r, "Year 5 invest users"); cx_usr = r
put_fm(ws, r, 3, f"=INDEX({S_ENG}!$C${t_usr}:$V${t_usr},1,20)", NUM); r += 1
label(ws, r, "Year 5 AUM per invest user (model)"); cx_per = r
put_fm(ws, r, 3, f"=INDEX({S_ENG}!$C${t_per}:$V${t_per},1,20)", M0); r += 1
label(ws, r, "Year 5 AUM per invest user REQUIRED", bold=True); cx_perreq = r
c = put_fm(ws, r, 3, f"=$C${cx_req}*G{cx_bal}/C{cx_usr}", M0, bold=True); c.fill = FILL_LT; r += 1
label(ws, r, "Multiple of modelled balance needed"); cx_mult = r
put_fm(ws, r, 3, f"=C{cx_perreq}/C{cx_per}", XF); r += 1
label(ws, r, "   for reference: Acorns avg balance"); put_in(ws, r, 3, 2142, M0)
ws.cell(r, 8, "mass-market micro-investing").font = F_NOTE; r += 1
label(ws, r, "   for reference: Betterment / Wealthfront avg"); put_in(ws, r, 3, 66000, M0)
ws.cell(r, 8, "primary-brokerage tier").font = F_NOTE; r += 2

section(ws, r, "  SENSITIVITY — REQUIRED AUM/BALANCE RATIO BY SLEEVE AND FEE", 8); r += 1
hdr(ws, r, 2, "AUM fee  \\  Vaulta sleeve %", "left")
sleeves = [0.10, 0.20, 0.35, 0.50, 0.58]
for i, s in enumerate(sleeves):
    c = hdr(ws, r, 3 + i, f"{s:.0%}")
grid_hdr = r; r += 1
for f in [0.0025, 0.0050, 0.0075, 0.0100]:
    label(ws, r, f"{f:.2%}")
    ws.cell(r, 2).font = F_IN
    for i, s in enumerate(sleeves):
        c = put_fm(ws, r, 3 + i, f"={TB}*{s}/{f}", XF)
        if abs(f - 0.005) < 1e-9 and abs(s - 0.20) < 1e-9:
            c.fill = FILL_LT; c.font = F_FMB
    r += 1
note(ws, r, "  Lower sleeve makes the crossover EASIER: at a 20% sleeve, AUM must reach 1.92x wallet balances; at 58% it must reach 5.57x."); r += 1
note(ws, r, "  Highlighted cell is the current setting (0.50% fee, 20% sleeve)."); r += 2

section(ws, r, "  READ-THROUGH", 8); r += 1
for t in [
  "The crossover does not depend on user counts. It is a pure ratio: AUM must exceed (T-bill rate x sleeve) / AUM fee times wallet balances.",
  "At the researched contribution rates (1.0-3.5% of balance per month) the model cannot reach that ratio inside five years. Contributions alone move AUM/balances by roughly adoption x contribution rate x months.",
  "The lever that actually moves it is the initial rollover — money transferred in from an existing bank or brokerage at signup — not the monthly drip.",
  "Raising the AUM fee is the other lever, and it is cheap: at 0.50% Vaulta undercuts Acorns' effective rate on a typical balance by roughly 3x, so there is headroom.",
]:
    c = ws.cell(r, 2, "• " + t); c.font = Font(size=9, color="404040")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 26
    r += 1

# ═══════════════════════════════════════════════════ SEGMENT UNIT ECONOMICS
ws = ws_ue
widths(ws, {'A': 3, 'B': 40, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 4, 'H': 15, 'I': 15})
title(ws, 1, "  SEGMENT UNIT ECONOMICS  —  at the 20/80 sleeve, with AUM", 9)
note(ws, 2, "  Steady-state view: each segment's AUM per investor is taken from the Year 5 engine output, so the AUM fee reflects accumulated balances rather than a guess.")
r = 4
SC = [3, 4, 5, 6]
hdr(ws, r, 2, "Metric", "left")
for i, nm in enumerate(seg_names): hdr(ws, r, SC[i], nm)
hdr(ws, r, 8, "Blended Y1"); hdr(ws, r, 9, "Blended Y3")
r += 1
def U(row, i): return f"{get_column_letter(SC[i])}{row}"
def urow(rlabel, fn, fmt, bold=False, fill=None):
    global r
    label(ws, r, rlabel, bold=bold)
    for i in range(4):
        c = put_fm(ws, r, SC[i], fn(i), fmt, bold=bold)
        if fill: c.fill = fill
    this = r; r += 1; return this

u_cac = urow("CAC", lambda i: f"={S_ASM}!$C${R['seg0']+i}", M2)
u_bal = urow("Wallet balance", lambda i: f"={S_ASM}!$D${R['seg0']+i}", M0)
u_chn = urow("Monthly churn", lambda i: f"={S_ASM}!$E${R['seg0']+i}", PCT2)
u_adp = urow("AUM adoption ceiling", lambda i: f"={S_ASM}!$F${R['seg0']+i}", PCT1)
u_aum = urow("AUM per investor (Yr 5, from engine)",
             lambda i: f"=INDEX({S_ENG}!$C${SEG_ROWS[i]['per']}:$V${SEG_ROWS[i]['per']},1,20)", M0)
r += 1
section(ws, r, "  MONTHLY REVENUE PER CONSUMER", 9); r += 1
u_nim = urow("T-bill NIM (20% sleeve)", lambda i: f"={U(u_bal,i)}*{TB}*{SLV}/12", M2)
u_spd = urow("   memo: monthly spend", lambda i: f"={U(u_bal,i)}*{SPND}", M2)
u_qr  = urow("QR merchant fee", lambda i: f"={U(u_spd,i)}*{QRS}*{QRF}", M2)
u_rin = urow("Rain interchange (net of txn cost)",
             lambda i: f"={U(u_spd,i)}*{CDS}*{RIN}-{U(u_spd,i)}*{CDS}/{TKT}*{RCT}", M2)
u_afe = urow("AUM fee", lambda i: f"={U(u_aum,i)}*{U(u_adp,i)}*{AUMF}/12", M2)
u_inf = urow("Infra", lambda i: f"=-{INFR}", M2)
u_net = urow("NET CONTRIBUTION / month",
             lambda i: f"={U(u_nim,i)}+{U(u_qr,i)}+{U(u_rin,i)}+{U(u_afe,i)}+{U(u_inf,i)}",
             M2, bold=True, fill=FILL_LT)
for i in range(4): ws.cell(u_net, SC[i]).border = TOPL
r += 1
section(ws, r, "  LIFETIME VALUE", 9); r += 1
u_d   = urow("   memo: decay factor", lambda i: f"=(1-{U(u_chn,i)})/(1+{MDSC})", '0.0000')
u_ltv = urow("Discounted LTV",
             lambda i: f"={U(u_net,i)}/(1+{MDSC})*(1-{U(u_d,i)}^{HORZ})/(1-{U(u_d,i)})",
             M0, bold=True, fill=FILL_LT)
u_rat = urow("LTV : CAC", lambda i: f"={U(u_ltv,i)}/{U(u_cac,i)}", XF, bold=True, fill=FILL_LT)
u_pbk = urow("CAC payback (months)",
             lambda i: f'=IF({U(u_net,i)}<=0,"n/a",{U(u_cac,i)}/{U(u_net,i)})', '0.0')
u_shr = urow("AUM share of contribution",
             lambda i: f"={U(u_afe,i)}/{U(u_net,i)}", PCT1)
r += 1

def blend2(row, ycol):
    mixcol = get_column_letter(2 + ycol)
    return "=" + " + ".join(f"{U(row,i)}*{S_ASM}!${mixcol}${R['mix0']+i}" for i in range(4))
for row, fmt in [(u_cac, M2), (u_bal, M0), (u_chn, PCT2), (u_adp, PCT1), (u_aum, M0),
                 (u_nim, M2), (u_spd, M2), (u_qr, M2), (u_rin, M2), (u_afe, M2),
                 (u_inf, M2), (u_net, M2), (u_ltv, M0)]:
    put_fm(ws, row, 8, blend2(row, 1), fmt)
    put_fm(ws, row, 9, blend2(row, 3), fmt)
put_fm(ws, u_rat, 8, f"=H{u_ltv}/H{u_cac}", XF, bold=True)
put_fm(ws, u_rat, 9, f"=I{u_ltv}/I{u_cac}", XF, bold=True)
put_fm(ws, u_pbk, 8, f"=H{u_cac}/H{u_net}", '0.0')
put_fm(ws, u_pbk, 9, f"=I{u_cac}/I{u_net}", '0.0')
put_fm(ws, u_shr, 8, f"=H{u_afe}/H{u_net}", PCT1)
put_fm(ws, u_shr, 9, f"=I{u_afe}/I{u_net}", PCT1)
for rr in (u_net, u_ltv, u_rat):
    ws.cell(rr, 8).fill = FILL_LT; ws.cell(rr, 9).fill = FILL_LT
r += 1
section(ws, r, "  SLEEVE SENSITIVITY — BLENDED Y3 LTV:CAC", 9); r += 1
hdr(ws, r, 2, "Vaulta sleeve %", "left")
for i, s in enumerate([0.10, 0.20, 0.35, 0.50, 0.58]): hdr(ws, r, 3 + i, f"{s:.0%}")
r += 1
label(ws, r, "Blended Y3 net contribution / month")
for i, s in enumerate([0.10, 0.20, 0.35, 0.50, 0.58]):
    terms = " + ".join(
        f"({U(u_bal,j)}*{TB}*{s}/12 + {U(u_qr,j)} + {U(u_rin,j)} + {U(u_afe,j)} + {U(u_inf,j)})"
        f"*{S_ASM}!$E${R['mix0']+j}" for j in range(4))
    c = put_fm(ws, r, 3 + i, "=" + terms, M2)
    if abs(s - 0.20) < 1e-9: c.fill = FILL_LT
r += 1
note(ws, r, "  The sleeve is the single biggest driver of consumer unit economics. Moving from 58% to 20% hands the consumer a 3.84% effective reward rate but cuts Vaulta's per-consumer NIM by about two thirds."); r += 1

# ═══════════════════════════════════════════════════ SENSITIVITY
ws = ws_sen
widths(ws, {'A': 3, 'B': 40, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 15, 'H': 15})
title(ws, 1, "  SENSITIVITY  —  Year 5 AUM fee revenue", 8)
note(ws, 2, "  Both grids recompute from the Assumptions sheet. They vary one driver at a time around the researched defaults.")
r = 4
section(ws, r, "  YEAR 5 AUM / WALLET-BALANCE RATIO  vs  rollover multiple and contribution rate", 8); r += 1
note(ws, r, "  Approximation: adoption x (rollover + contribution rate x months invested). Months invested is weighted by the availability ramp."); r += 1
hdr(ws, r, 2, "contribution %/mo  \\  rollover x", "left")
rolls = [0.0, 0.5, 1.0, 2.0, 4.0, 8.0]
for i, rv in enumerate(rolls): hdr(ws, r, 3 + i, f"{rv:.1f}x")
r += 1
ADP = f"INDEX({S_ENG}!$C${t_adp}:$V${t_adp},1,20)"
MONTHS = 36
for cr in [0.01, 0.02, 0.03, 0.05, 0.08]:
    label(ws, r, f"{cr:.0%}")
    ws.cell(r, 2).font = F_IN
    for i, rv in enumerate(rolls):
        c = put_fm(ws, r, 3 + i, f"={ADP}*({rv}+{cr}*{MONTHS})", XF)
        if abs(cr - 0.02) < 1e-9 and abs(rv - 1.0) < 1e-9:
            c.fill = FILL_LT; c.font = F_FMB
    r += 1
label(ws, r, "Required ratio to overtake NIM", bold=True)
c = put_fm(ws, r, 3, f"={S_CX}!$C${cx_req}", XF, bold=True); c.fill = FILL_WARN
r += 2
note(ws, r, "  Cells at or above the required ratio are combinations where AUM fee revenue exceeds T-bill NIM in Year 5."); r += 2

section(ws, r, "  YEAR 5 AUM FEE REVENUE  vs  AUM fee rate", 8); r += 1
hdr(ws, r, 2, "AUM fee", "left")
for i, f in enumerate([0.0025, 0.0050, 0.0075, 0.0100, 0.0150]): hdr(ws, r, 3 + i, f"{f:.2%}")
r += 1
label(ws, r, "Year 5 AUM fee revenue")
for i, f in enumerate([0.0025, 0.0050, 0.0075, 0.0100, 0.0150]):
    c = put_fm(ws, r, 3 + i, f"={S_REV}!$G${an_aumbal}*{f}", M0)
    if abs(f - 0.005) < 1e-9: c.fill = FILL_LT
r += 1
label(ws, r, "vs Year 5 T-bill NIM")
for i in range(5): put_fm(ws, r, 3 + i, f"={S_REV}!$G${an_nim}", M0)
r += 1
label(ws, r, "AUM / NIM", bold=True)
for i in range(5):
    cl = get_column_letter(3 + i)
    put_fm(ws, r, 3 + i, f"={cl}{r-2}/{cl}{r-1}", XF, bold=True)
r += 2
note(ws, r, "  Acorns' flat $3/month is roughly 1.7% of its average balance, so even 1.50% here would undercut the closest mass-market comparable.")

# ═══════════════════════════════════════════════════ COVER
ws = ws_cov
widths(ws, {'A': 3, 'B': 52, 'C': 18, 'D': 18, 'E': 18, 'F': 18, 'G': 18})
title(ws, 1, "  VAULTA — AUM MODEL v1", 7)
c = ws.cell(2, 1, "  Does tokenized-investment AUM overtake the T-bill sleeve in Years 3-5?")
c.font = Font(color="FFFFFF", size=10, italic=True)
for cc in range(1, 8): ws.cell(2, cc).fill = FILL_TITLE
r = 4

section(ws, r, "  THE MODEL", 7); r += 1
note(ws, r, "  Customers x adoption x monthly contribution (% of wallet balance) x initial rollover at signup, compounding at the expected"); r += 1
note(ws, r, "  investment return, net of redemptions and churn leakage, charged at the AUM fee. Levers are per segment on the Assumptions sheet."); r += 2

section(ws, r, "  THE ANSWER", 7); r += 1
label(ws, r, "Required AUM as a multiple of wallet balances", bold=True)
c = put_fm(ws, r, 3, f"={S_CX}!$C${cx_req}", XF, bold=True); c.fill = FILL_LT; r += 1
label(ws, r, "Year 5 actual AUM / balances")
put_fm(ws, r, 3, f"={S_CX}!$G${cx_act}", XF, bold=True); r += 1
label(ws, r, "Shortfall", bold=True)
c = put_fm(ws, r, 3, f"={S_CX}!$G${cx_gap}", XF, bold=True); c.fill = FILL_WARN; r += 1
label(ws, r, "Does AUM overtake NIM by Year 5?", bold=True)
c = put_fm(ws, r, 3, f"={S_REV}!$G${an_ovr}", None, bold=True); c.fill = FILL_WARN; r += 2

section(ws, r, "  ANNUAL REVENUE — WHAT VAULTA KEEPS", 7); r += 1
hdr(ws, r, 2, "Metric", "left")
for i in range(5): hdr(ws, r, 3 + i, f"Year {i+1}")
r += 1
for nm, srow, fmt, bold in [("T-bill NIM (20% sleeve)", an_nim, M0, True),
                            ("Transaction fees", an_fee, M0, False),
                            ("AUM fee", an_aum, M0, True),
                            ("Total retained revenue", an_tot, M0, True),
                            ("AUM as % of retained revenue", an_pct, PCT1, False),
                            ("AUM / NIM ratio", an_rat, XF, True),
                            ("Closing AUM", an_aumbal, M0, False),
                            ("Total wallet balances", an_wal, M0, False)]:
    label(ws, r, nm, bold=bold)
    for i in range(5):
        cl = get_column_letter(3 + i)
        cc = put_fm(ws, r, 3 + i, f"={S_REV}!${cl}${srow}", fmt, bold=bold)
        if nm == "AUM / NIM ratio": cc.fill = FILL_LT
    r += 1
r += 1

section(ws, r, "  UNIT ECONOMICS AT THE 20/80 SLEEVE", 7); r += 1
hdr(ws, r, 2, "Metric", "left"); hdr(ws, r, 3, "Year 1 mix"); hdr(ws, r, 4, "Year 3 mix")
r += 1
for nm, srow, fmt in [("Blended CAC", u_cac, M2), ("Blended net contribution / month", u_net, M2),
                      ("Blended discounted LTV", u_ltv, M0), ("Blended LTV : CAC", u_rat, XF),
                      ("AUM share of contribution", u_shr, PCT1)]:
    label(ws, r, nm)
    put_fm(ws, r, 3, f"={S_UE}!$H${srow}", fmt)
    put_fm(ws, r, 4, f"={S_UE}!$I${srow}", fmt)
    r += 1
r += 1

section(ws, r, "  FINDINGS", 7); r += 1
for t, red in [
 ("AUM does NOT overtake the T-bill sleeve by Year 5 on researched contribution rates. The crossover needs AUM to reach 1.92x total wallet balances; the model reaches a small fraction of that.", True),
 ("The crossover is a pure ratio and does not depend on how many customers Vaulta has: AUM must exceed (T-bill rate x sleeve) / AUM fee, times wallet balances. Growth alone never triggers it.", False),
 ("Moving the sleeve to 20/80 makes the crossover EASIER, not harder — it cuts the NIM hurdle from 5.57x balances to 1.92x. But it does so by shrinking the NIM line, not by growing AUM.", False),
 ("COST OF THE 20/80 SPLIT: blended LTV:CAC falls to 1.47x at the Year 3 mix, against a 3x investability bar. Segment 1 (Underbanked) turns value-destructive at 0.32x and Segment 2 only reaches 1.21x. Only Segments 3 and 4 clear their CAC.", True),
 ("The monthly drip is not the lever. At 1-3.5% of balance per month, contributions move AUM/balances by roughly adoption x rate x months — far too slowly. The initial rollover at signup is what moves it.", False),
 ("The 0.50% fee is conservatively priced. Acorns' flat $3/month is about 1.7% of its average balance, so there is real headroom before Vaulta looks expensive.", False),
 ("Redemption/leakage at 8%/yr is a modelling judgement, not a researched figure. It is the weakest default in the workbook and is flagged on the Research sheet.", True),
]:
    c = ws.cell(r, 2, "• " + t)
    c.font = F_RED if red else Font(size=9, color="404040")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    r += 1

for s in wb.worksheets:
    s.sheet_view.showGridLines = False

os.makedirs(REPO, exist_ok=True)
wb.save(OUT)
print("WROTE", OUT)
